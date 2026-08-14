"""「目录助手」面板：从 DWG/DXF 图纸按图纸模板取值，生成图纸目录 Excel。

- 待处理区：图纸文件多选列表（选择/追加/右键删除/Delete/拖放）
- 数据源区：图纸模板库下拉（上传/删除到 templates/catalog）+ 表格模板（Excel）
- 输出区：输出目录（默认 = 第一个图纸文件所在目录 output）
- 运行区：开始/停止 + 进度条；日志区

后台线程执行 catalog_pipeline（模板解析 → 取值 → 文件粒度目录），
日志与进度经队列回传。模板库与填表助手分离（templates/catalog vs templates/fill）。
"""

from __future__ import annotations

import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

from cadbatchassistant.core.app_config import (
    get_oda,
    get_out_version,
    load_catalog_rules,
)
from cadbatchassistant.core.catalog_excel_writer import detect_sheet_candidates
from cadbatchassistant.core.catalog_pipeline import (
    PipelineResult,
    parse_template_anchors,
    run_pipeline,
)
from cadbatchassistant.core.catalog_template_reader import collect_fields
from cadbatchassistant.core.templates import templates_dir
from cadbatchassistant.gui.async_panel import AsyncPanel
from cadbatchassistant.gui.gui_shared import (
    FilesPanelMixin,
    PanelLayoutMixin,
    RunStartMixin,
    TemplateLibraryMixin,
    finish_popup,
    load_panel_config,
    save_panel_config,
    warn_require,
)
from cadbatchassistant.gui.tk_util import center_window


class CatalogPanel(
    FilesPanelMixin, TemplateLibraryMixin, PanelLayoutMixin, RunStartMixin, AsyncPanel
):
    """目录生成面板（模板标记取值）；文件列表/模板库复用共享组件。"""

    TEMPLATE_CATEGORY = "catalog"
    TEMPLATE_CONFIG_KEY = "catalog_template"
    TEMPLATE_UPLOAD_TITLE = "上传图纸模板（[字段名] 取值位置）"

    def __init__(self, parent: tk.Widget) -> None:
        super().__init__(parent)
        self.scanned_files: list[str] = []
        self._last_result: PipelineResult | None = None
        self._build_ui()
        self._load()

    # ---------------- UI ----------------
    def _build_ui(self) -> None:
        self._main = self._build_root()

        # 1. 待处理区：图纸文件多选
        self._add_input_section(width=60, bind_delete=True)

        # 2. 数据源区：表格模板 + 图纸模板
        self._add_src_section(
            "表格模板:",
            (".xlsx",),
            on_xlsx_hit=lambda h: save_panel_config({"catalog_xlsx": h}),
            tpl_width=24,
        )

        # 3. 输出区
        self.var_out = tk.StringVar()
        self._add_output_section(self.var_out)

        # 4. 运行区（ODA 路径与输出版本在「设置」tab，全局共享）
        self._add_run_section(maximum=100)

        # 5. 日志区
        self._add_log_section()

    # ---------------- 表格模板（Excel） ----------------
    def _browse_xlsx(self) -> None:
        f = filedialog.askopenfilename(
            title="选择表格模板",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if f:
            self.var_xlsx.set(f)
            save_panel_config({"catalog_xlsx": f})

    # ---------------- 输出目录 ----------------
    def _default_output(self) -> None:
        super()._default_output()
        if self.var_out.get().strip():
            save_panel_config({"catalog_out": self.var_out.get()})

    def _browse_dir(self, var: tk.StringVar) -> None:
        super()._browse_dir(var)
        if var.get().strip():
            save_panel_config({"catalog_out": var.get()})

    def _on_drop_out_dir(self, event) -> None:
        super()._on_drop_out_dir(event)
        if self.var_out.get().strip():
            save_panel_config({"catalog_out": self.var_out.get()})

    # ---------------- 配置记忆 ----------------
    def _load(self) -> None:
        self._refresh_templates()  # 恢复上次选择的图纸模板
        cfg = load_panel_config()
        last_xlsx = cfg.get("catalog_xlsx", "")
        if last_xlsx and os.path.isfile(last_xlsx):
            self.var_xlsx.set(last_xlsx)
        last_out = cfg.get("catalog_out", "")
        if last_out and os.path.isdir(last_out):
            self.var_out.set(last_out)

    # ---------------- 运行 ----------------
    def _prepare_run(self) -> tuple | None:
        """校验输入 + 预检（模板解析/sheet 定位）并收集 worker 参数。

        表格模板 sheet 预检在主线程完成：解析模板锚点（字段名 + 取值区域）
        并反推 sheet，无匹配提前弹错；多个 sheet 并列最高时让用户选择。
        锚点随后传入 pipeline 复用，省一次模板 DXF 转换。
        返回 (template, xlsx, files, out, oda, version, rules, sheet_name,
        anchors)；校验/预检失败返回 None（已弹窗）。
        """
        tpl_name = self.var_template.get().strip()
        template = str(templates_dir("catalog") / tpl_name) if tpl_name else ""
        xlsx = self.var_xlsx.get().strip()
        files = list(self.scanned_files)
        out = self.var_out.get().strip()

        if not warn_require(
            bool(tpl_name) and os.path.isfile(template),
            "请从图纸模板下拉框选择模板（可先「上传」）",
        ):
            return None
        if not warn_require(
            bool(xlsx) and os.path.isfile(xlsx), "请选择有效的表格模板"
        ):
            return None
        if not warn_require(bool(files), "请选择要处理的 DWG/DXF 文件"):
            return None
        if not out:
            self._default_output()
            out = self.var_out.get().strip()
        if not warn_require(bool(out), "请设置输出目录"):
            return None

        oda = get_oda()
        try:
            anchors = parse_template_anchors(template, oda)
        except Exception as ex:  # noqa: BLE001 - 模板解析失败提前提示
            messagebox.showerror("图纸目录助手", f"模板解析失败：{ex}")
            return None
        fields = collect_fields(anchors)
        sheet_name: str | None = None
        try:
            wb = load_workbook(xlsx)
            cands = detect_sheet_candidates(wb, fields)
            wb.close()
        except Exception as ex:  # noqa: BLE001 - 表格模板读取失败提前提示
            messagebox.showerror("图纸目录助手", f"读取表格模板失败：{ex}")
            return None
        if not cands:
            messagebox.showerror(
                "图纸目录助手",
                "表格模板中未找到与字段匹配的表头（sheet 与表头行）（字段："
                + "、".join(fields)
                + "）。表头列名应包含与图纸模板 [字段名] 占位符一致的字段名。",
            )
            return None
        tied = [c for c in cands if c[0] == cands[0][0]]
        if len(tied) > 1:
            sheet_name = self._ask_sheet(tied, fields)
            if sheet_name is None:
                return None
        return (
            template,
            xlsx,
            files,
            out,
            oda,
            get_out_version(),
            load_catalog_rules(),
            sheet_name,
            anchors,
        )

    def _after_begin_run(self, args: tuple) -> None:
        """begin_run 之后输出 sheet 定位日志（worker 参数 args[7] 为 sheet 名）。"""
        sheet_name = args[7]
        self._emit(f"表格模板 sheet：{sheet_name or '自动定位'}")

    def _ask_sheet(self, candidates, fields: list[str]) -> str | None:
        """多个 sheet 并列最高时弹窗选择，返回所选 sheet 名；取消/关闭返回 None。"""
        names = [c[1].title for c in candidates]
        pick: dict[str, str | None] = {"name": None}
        win = tk.Toplevel(self._root)
        win.title("选择表格模板 sheet")
        win.transient(self._root)
        win.resizable(False, False)
        ttk.Label(
            win,
            text="表格模板中有多个 sheet 与字段匹配（字段："
            + "、".join(fields)
            + "），请选择本次使用的 sheet：",
            wraplength=420,
        ).pack(padx=12, pady=(12, 6))
        lb = tk.Listbox(win, height=min(len(names), 8), width=40, selectmode="single")
        for n in names:
            lb.insert("end", n)
        lb.selection_set(0)
        lb.pack(padx=12, pady=4)

        def _ok() -> None:
            sel = lb.curselection()
            if sel:
                pick["name"] = names[sel[0]]
            win.destroy()

        def _cancel() -> None:
            win.destroy()

        btns = ttk.Frame(win)
        btns.pack(pady=(4, 12))
        ttk.Button(btns, text="确定", command=_ok).pack(side="left", padx=6)
        ttk.Button(btns, text="取消", command=_cancel).pack(side="left", padx=6)
        win.protocol("WM_DELETE_WINDOW", _cancel)
        center_window(win, self._root)  # 相对主窗口居中
        win.grab_set()
        self._root.wait_window(win)
        return pick["name"]

    def _work(
        self, template, xlsx, files, out, oda, version, rules, sheet_name, anchors=None
    ) -> bool:
        res = run_pipeline(
            template,
            xlsx,
            files,
            out,
            oda,
            version,
            rules,
            sheet_name=sheet_name,
            log=self._emit,
            progress=lambda p: self._emit(None, int(p)),
            is_cancelled=self._is_cancelled,
            template_anchors=anchors,
        )
        self._last_result = res
        if res.error:
            self._emit(f"处理失败：{res.error}")
        return res.ok

    def _finish_notify(self, success: bool) -> None:
        """目录助手专属完成弹窗（覆盖 PanelLayoutMixin 默认的 finish_popup）。

        按钮复位由 AsyncPanel._on_finish 统一完成，这里只负责统计弹窗；
        弹窗标题与改字/填表助手统一为「完成」（成功 showinfo / 失败
        showwarning），正文保留目录统计信息。
        """
        res = self._last_result
        if res is None:
            finish_popup(success)
            return
        if success and res.out_path:
            msg = (
                f"目录生成完成：{res.out_path}\n\n"
                f"图纸：{res.total_files} 个文件\n"
                f"无管段(NA)：{res.na_rows} 张\n"
                f"总页数：{res.total_pages}\n"
                f"字段：{' / '.join(res.fields)}\n"
            )
            if res.failed_files:
                msg += f"\n转换失败：{len(res.failed_files)} 个\n" + "\n".join(
                    res.failed_files[:10]
                )
            messagebox.showinfo("完成", msg)
        elif not success:
            messagebox.showwarning("完成", f"处理失败：\n{res.error or '未知错误'}")
