# -*- coding: utf-8 -*-
"""模板占位扫描（learn_spec）测试：sheet 透传与按名称精确匹配。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import ezdxf
import openpyxl

from cadbatchassistant.core import fill_learn_spec


def _make_xlsx_multi(path, sheets: dict) -> None:
    """写多工作表临时 xlsx：{工作表名: (表头, 行列表)}。"""
    wb = openpyxl.Workbook()
    first = True
    for name, (header, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(header)
        for r in rows:
            ws.append(r)
    wb.save(path)
    wb.close()


def _make_dxf_with_placeholders(path, texts: list[str],
                                layer: str = "0") -> None:
    """写含占位符 TEXT 的临时 DXF（默认图层 0，可指定图层）。"""
    doc = ezdxf.new("R2004")
    msp = doc.modelspace()
    for i, t in enumerate(texts):
        msp.add_text(t, dxfattribs={
            "insert": (10, 10 + i * 10), "height": 3.0, "layer": layer})
    doc.saveas(path)


class ScanPlaceholdersSheetTest(unittest.TestCase):
    def setUp(self):
        self._tmp = Path(tempfile.mkdtemp(prefix="learn_sheet_"))

    def tearDown(self):
        for f in self._tmp.iterdir():
            f.unlink(missing_ok=True)
        self._tmp.rmdir()

    def test_sheet_affects_header_source(self):
        xlsx = self._tmp / "data.xlsx"
        _make_xlsx_multi(xlsx, {
            "SheetA": (["图号", "名称"], [["A-1", "图1"]]),
            "SheetB": (["图号", "压力"], [["A-1", "1.5"]]),
        })
        dxf = self._tmp / "t.dxf"
        _make_dxf_with_placeholders(dxf, ["[压力]"])

        # 默认第一个 sheet（SheetA 无「压力」列）→ 匹配不到（spec 为空）
        by_default = fill_learn_spec.scan_placeholders(str(dxf), str(xlsx))
        self.assertEqual(by_default, {})

        # 指定 SheetB → 按名称精确匹配到「压力」
        by_sheet = fill_learn_spec.scan_placeholders(
            str(dxf), str(xlsx), sheet="SheetB")
        self.assertEqual(list(by_sheet["0"].keys()), ["压力"])

    def test_exact_match_no_normalization(self):
        xlsx = self._tmp / "data.xlsx"
        _make_xlsx_multi(xlsx, {"数据表": (["Name", "图号"], [["x", "A-1"]])})
        dxf = self._tmp / "t.dxf"
        _make_dxf_with_placeholders(dxf, ["[name]", "[图号]"])

        spec = fill_learn_spec.scan_placeholders(str(dxf), str(xlsx))
        # 「name」≠「Name」不匹配（不归一化）；「图号」精确匹配
        self.assertEqual(list(spec["0"].keys()), ["图号"])

    def test_placeholder_in_non_zero_layer(self):
        # 占位符位于非 "0" 图层也能扫描匹配（任意图层）
        xlsx = self._tmp / "data.xlsx"
        _make_xlsx_multi(xlsx, {"数据表": (["图号", "名称"], [["A-1", "图1"]])})
        dxf = self._tmp / "t.dxf"
        _make_dxf_with_placeholders(dxf, ["[图号]", "[名称]"], layer="TEXT1")

        spec = fill_learn_spec.scan_placeholders(str(dxf), str(xlsx))
        self.assertEqual(list(spec["TEXT1"].keys()), ["图号", "名称"])

    def test_placeholders_in_multiple_layers_grouped(self):
        # 多个图层同时存在占位符时按图层分组返回
        xlsx = self._tmp / "data.xlsx"
        _make_xlsx_multi(xlsx, {"数据表": (["图号", "名称"], [["A-1", "图1"]])})
        dxf = self._tmp / "t.dxf"
        _make_dxf_with_placeholders(dxf, ["[图号]"], layer="0")
        doc = ezdxf.readfile(dxf)
        doc.modelspace().add_text(
            "[名称]", dxfattribs={
                "insert": (50, 10), "height": 3.0, "layer": "TEXT1"})
        doc.saveas(dxf)

        spec = fill_learn_spec.scan_placeholders(str(dxf), str(xlsx))
        self.assertEqual(list(spec["0"].keys()), ["图号"])
        self.assertEqual(list(spec["TEXT1"].keys()), ["名称"])


if __name__ == "__main__":
    unittest.main()
