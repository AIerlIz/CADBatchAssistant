"""fill_learn_spec 扫描与填表 pipeline meta 路径单测。

- scan_all_placeholders：扫描全部 [列名] 占位符（不依赖数据表）、
  entity_desc 可 JSON 序列化（Vec3 已清理）
- value_rule_for：当前恒为 ('value', '')
- 模板 meta 往返：scan_all_placeholders → save/load_template_meta 一致
- fill_pipeline meta 路径：读 meta 生成规格并填表（entity desc dict 重建），
  且不再调用 template_to_dxf
- fill_pipeline CLI 兜底：模板无 meta 时现场转换扫描（template_to_dxf 被调用）
"""

from __future__ import annotations

import json
from unittest import mock

import ezdxf
import openpyxl

from cadbatchassistant.core.common.template_meta import (
    load_template_meta,
    save_template_meta,
)
from cadbatchassistant.core.fill.fill_learn_spec import (
    scan_all_placeholders,
    value_rule_for,
)


def _make_tpl_dxf(path, placeholders=("[图号]", "[名称]")):
    doc = ezdxf.new("R2004")
    m = doc.modelspace()
    for i, t in enumerate(placeholders):
        m.add_text(
            t,
            dxfattribs={
                "insert": (10, 10 + i * 10),
                "height": 3.0,
                "layer": "T1",
            },
        )
    m.add_text("普通文字", dxfattribs={"insert": (0, 0), "height": 1.0})
    doc.saveas(path)
    return path


def _make_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据表"
    ws.append(rows[0])
    for r in rows[1:]:
        ws.append(r)
    wb.save(path)
    return path


def _make_before_dxf(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = ezdxf.new("R2004")
    doc.modelspace().add_line((0, 0), (30, 30))
    doc.saveas(path)
    return path


def test_scan_all_placeholders_collects_all(tmp_path):
    p = _make_tpl_dxf(tmp_path / "t.dxf")
    phs = scan_all_placeholders(str(p))
    assert len(phs) == 2
    assert {x["text"] for x in phs} == {"图号", "名称"}
    assert all(x["layer"] == "T1" for x in phs)
    assert all(x["height"] == 3.0 for x in phs)


def test_scan_all_placeholders_json_serializable(tmp_path):
    p = _make_tpl_dxf(tmp_path / "t.dxf")
    phs = scan_all_placeholders(str(p))
    data = json.dumps(phs, ensure_ascii=False)  # 不抛
    back = json.loads(data)
    # Vec3 → 可 JSON 的基础类型（insert 从 Vec3 转为 list）
    assert isinstance(back[0]["entity_desc"]["attribs"]["insert"], list)


def test_value_rule_for_constant():
    assert value_rule_for("NPD (inch)") == ("value", "")
    assert value_rule_for("设计压力") == ("value", "")


def test_meta_roundtrip_placeholders(tmp_path):
    tpl = _make_tpl_dxf(tmp_path / "tpl.dxf")
    phs = scan_all_placeholders(str(tpl))
    save_template_meta(tpl, {"placeholders": phs})
    meta = load_template_meta(tpl)
    assert meta is not None
    assert meta["placeholders"] == phs


def test_fill_pipeline_meta_path(tmp_path):
    """完整 meta 路径：读模板 meta（不转换模板）→ 按表头匹配 → 填表。"""
    from cadbatchassistant.core.fill.fill_pipeline import run_pipeline

    tpl = _make_tpl_dxf(tmp_path / "tpl.dxf", placeholders=("[图号]",))
    save_template_meta(tpl, {"placeholders": scan_all_placeholders(str(tpl))})
    xlsx = _make_xlsx(tmp_path / "data.xlsx", [["图纸名", "图号"], ["A1", "ABC-001"]])
    _make_before_dxf(tmp_path / "in" / "A1.dxf")
    out = tmp_path / "out"
    conv = mock.Mock()
    conv.resolve.return_value = ""
    conv.require_for_dwg.return_value = None
    with mock.patch(
        "cadbatchassistant.core.fill.fill_pipeline.dc.get_converter", return_value=conv
    ):
        summary = run_pipeline(
            str(xlsx),
            str(tmp_path / "in"),
            str(out),
            template=str(tpl),
            inputs=["A1"],
            emit=lambda m: None,
        )
    assert summary["ok"] == 1, summary
    filled = ezdxf.readfile(str(out / "A1.dxf"))
    texts = [e.dxf.text for e in filled.modelspace() if e.dxftype() == "TEXT"]
    assert "ABC-001" in texts
    # meta 路径不应触发模板转换
    conv.template_to_dxf.assert_not_called()


def test_fill_pipeline_cli_fallback_without_meta(tmp_path):
    """CLI 兜底：模板无 meta 时现场转换 + 扫描（template_to_dxf 被调用）。"""
    from cadbatchassistant.core.fill.fill_pipeline import run_pipeline

    tpl = _make_tpl_dxf(tmp_path / "tpl.dxf", placeholders=("[图号]",))
    xlsx = _make_xlsx(tmp_path / "data.xlsx", [["图纸名", "图号"], ["A1", "ABC-001"]])
    _make_before_dxf(tmp_path / "in" / "A1.dxf")
    out = tmp_path / "out"
    conv = mock.Mock()
    conv.resolve.return_value = ""
    conv.require_for_dwg.return_value = None
    conv.template_to_dxf.return_value = str(tpl)  # 模板已是 DXF，直接返回
    with mock.patch(
        "cadbatchassistant.core.fill.fill_pipeline.dc.get_converter", return_value=conv
    ):
        summary = run_pipeline(
            str(xlsx),
            str(tmp_path / "in"),
            str(out),
            template=str(tpl),
            inputs=["A1"],
            emit=lambda m: None,
        )
    assert summary["ok"] == 1, summary
    conv.template_to_dxf.assert_called_once()
