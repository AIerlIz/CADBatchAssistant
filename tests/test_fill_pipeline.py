"""fill_pipeline 流程保护测试：输出目录与输入目录重合时拒绝处理，防止覆盖源文件。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from cadbatchassistant.core.fill import fill_pipeline


class OutputSameAsInputProtectionTest(unittest.TestCase):
    """H1：输出目录与输入目录（或源文件所在目录）重合时必须报错，不得删除/覆盖源文件。"""

    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp(prefix="fill_pipeline_test_")
        self.before_dir = Path(self.tmp) / "input"
        self.before_dir.mkdir()
        # 放一个真实 DXF 作为"源图纸"，用于验证重合时不会被触碰
        self.src_dxf = self.before_dir / "A1.dxf"
        self.src_dxf.write_text("MOCK-DXF-ORIGINAL", encoding="utf-8")
        self.xlsx = Path(self.tmp) / "data.xlsx"
        self.xlsx.write_text("mock", encoding="utf-8")

    def tearDown(self) -> None:
        import shutil

        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_run_pipeline_same_dir_rejected(self) -> None:
        """out_dir == before_dir 时抛 ValueError，且源文件原样保留。"""
        with self.assertRaises(ValueError) as ctx:
            fill_pipeline.run_pipeline(
                str(self.xlsx),
                str(self.before_dir),
                str(self.before_dir),
                template=str(self.src_dxf),
            )
        self.assertIn("不能与输入图纸目录相同", str(ctx.exception))
        # 源文件未被删除/覆盖
        self.assertEqual(self.src_dxf.read_text(encoding="utf-8"), "MOCK-DXF-ORIGINAL")

    def test_run_pipeline_files_same_source_dir_rejected(self) -> None:
        """run_pipeline_files：输出目录 == 源文件所在目录时抛 ValueError。"""
        with self.assertRaises(ValueError) as ctx:
            fill_pipeline.run_pipeline_files(
                str(self.xlsx), [str(self.src_dxf)], str(self.before_dir)
            )
        self.assertIn("不能与输入图纸所在目录相同", str(ctx.exception))
        self.assertEqual(self.src_dxf.read_text(encoding="utf-8"), "MOCK-DXF-ORIGINAL")

    def test_run_pipeline_different_dir_not_rejected(self) -> None:
        """输出目录与输入目录不同时，不应触发重合保护。

        用 mock 阻断模板转换（模拟 ODA 转换阶段失败），断言异常来自流程
        本身（ODAError）而非重合保护（ValueError"目录相同"）。
        """
        out = Path(self.tmp) / "output"
        from unittest import mock

        from cadbatchassistant.core import dwg_converter as dc

        conv = mock.Mock()
        conv.resolve.return_value = ""
        conv.require_for_dwg.return_value = None
        conv.template_to_dxf.side_effect = dc.ODAError("mock oda 失败")
        with (
            mock.patch.object(dc, "get_converter", return_value=conv),
            self.assertRaises(dc.ODAError),
        ):
            fill_pipeline.run_pipeline(
                str(self.xlsx),
                str(self.before_dir),
                str(out),
                template=str(self.src_dxf),
            )

    def test_run_pipeline_files_duplicate_basename_rejected(self) -> None:
        """M2：跨目录同名文件（大小写不敏感）复制到临时目录会互相覆盖 → 拒绝。"""
        src2 = Path(self.tmp) / "input2"
        src2.mkdir()
        dup = src2 / "A1.dxf"
        dup.write_text("MOCK-DXF-2", encoding="utf-8")
        with self.assertRaises(ValueError) as ctx:
            fill_pipeline.run_pipeline_files(
                str(self.xlsx),
                [str(self.src_dxf), str(dup)],
                str(Path(self.tmp) / "output"),
            )
        self.assertIn("输入文件重名", str(ctx.exception))

    def test_run_pipeline_files_duplicate_case_insensitive_rejected(self) -> None:
        """M2：大小写不同也视为重名（a.dwg vs A.DWG 会互相覆盖）。"""
        src2 = Path(self.tmp) / "input2"
        src2.mkdir()
        dup = src2 / "a1.DXF"  # 与 A1.dxf 大小写不同
        dup.write_text("MOCK-DXF-2", encoding="utf-8")
        with self.assertRaises(ValueError) as ctx:
            fill_pipeline.run_pipeline_files(
                str(self.xlsx),
                [str(self.src_dxf), str(dup)],
                str(Path(self.tmp) / "output"),
            )
        self.assertIn("输入文件重名", str(ctx.exception))

    def test_dwg_priority_when_same_stem_dwg_and_dxf(self) -> None:
        """M8：同名 .dwg 与 .dxf 共存时按 DWG 处理（dwg 优先，结果确定）。

        通过观测 [1/4] 阶段是否调用 dwg_to_dxf 来验证：
        输入目录同时有 A1.dwg 与 A1.dxf 时，dwg 分支被选中（dxf 分支只是 copy2）。
        """
        from unittest import mock

        from cadbatchassistant.core import dwg_converter as dc

        self.before_dir.joinpath("A1.dwg").write_bytes(b"MOCK-DWG")
        conv = mock.Mock()
        conv.resolve.return_value = ""
        conv.require_for_dwg.return_value = None
        conv.require_for_dwg.return_value = None
        conv.dwg_to_dxf.side_effect = dc.ODAError("mock oda 失败")
        with (
            mock.patch.object(
                dc,
                "get_converter",
                return_value=conv,
            ),
        ):
            with self.assertRaises(dc.ODAError):
                fill_pipeline.run_pipeline(
                    str(self.xlsx),
                    str(self.before_dir),
                    str(Path(self.tmp) / "output"),
                    template=str(self.src_dxf),
                )
            # 若被误判为 dxf，则不会调用 dwg_to_dxf（而只是 copy2 复制）
            conv.dwg_to_dxf.assert_called_once()

    def test_tmp_dir_cleaned_when_auto_created(self) -> None:
        """M3：workdir=None 时自建临时目录，流程结束（含异常路径）后应被清理。"""
        import tempfile
        from unittest import mock

        from cadbatchassistant.core import dwg_converter as dc

        before = list(Path(tempfile.gettempdir()).glob("iso_fill_*"))
        conv = mock.Mock()
        conv.resolve.return_value = ""
        conv.require_for_dwg.return_value = None
        conv.template_to_dxf.side_effect = dc.ODAError("mock oda 失败")
        with (
            mock.patch.object(dc, "get_converter", return_value=conv),
            self.assertRaises(dc.ODAError),
        ):
            fill_pipeline.run_pipeline(
                str(self.xlsx),
                str(self.before_dir),
                str(Path(self.tmp) / "output"),
                template=str(self.src_dxf),
            )
        after = list(Path(tempfile.gettempdir()).glob("iso_fill_*"))
        # 异常路径后不残留自建临时目录
        self.assertEqual(after, before)


class FillPipelineWriteBackSkipTest(unittest.TestCase):
    """B1 回归：fill 阶段 skipped 的图纸在 [4/4] 写回时被跳过，整批不崩。

    修复前 skip 集合（无扩展名 stem）与 write_back_dxf_batch 收到的文件名
    （含扩展名）不匹配，skipped 的 DXF 会在 copy2 时 FileNotFoundError
    中断整批；本用例走完整 run_pipeline（真实 xlsx/模板/before DXF，
    仅 mock ODA converter），验证 [4/4] 阶段正常产出且跳过图不写回。
    """

    def setUp(self) -> None:
        import ezdxf
        import openpyxl

        self.tmp = Path(tempfile.mkdtemp(prefix="fill_pipe_skip_"))
        self.before_dir = self.tmp / "input"
        self.out_dir = self.tmp / "output"
        self.before_dir.mkdir()
        self.out_dir.mkdir()
        # 两张 before DXF：A1（含匹配文字 D-001，在数据表）、B1（不在数据表）
        import ezdxf as _ezdxf
        doc_a = _ezdxf.new("R2013")
        doc_a.modelspace().add_text(
            "D-001", dxfattribs={"insert": (10, 10, 0), "height": 3.0}
        )
        doc_a.saveas(self.before_dir / "A1.dxf")
        doc_b = _ezdxf.new("R2013")
        doc_b.saveas(self.before_dir / "B1.dxf")
        # 模板 DXF：占位符 [图号] 与数据表表头匹配
        self.template = self.tmp / "template.dxf"
        doc = ezdxf.new("R2013")
        doc.modelspace().add_text(
            "[图号]", dxfattribs={"insert": (10, 20, 0), "height": 3.0}
        )
        doc.saveas(self.template)
        # xlsx 只含 A1
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["图纸", "图号"])
        ws.append(["A1", "D-001"])
        self.xlsx = self.tmp / "data.xlsx"
        wb.save(self.xlsx)
        wb.close()

    def tearDown(self) -> None:
        import shutil

        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_skipped_dxf_not_written_back(self) -> None:
        from unittest import mock

        from cadbatchassistant.core import dwg_converter as dc
        from cadbatchassistant.core.fill import fill_pipeline

        conv = mock.Mock()
        conv.resolve.return_value = ""
        conv.require_for_dwg.return_value = None
        conv.require_for_dwg.return_value = None
        conv.template_to_dxf.return_value = str(self.template)
        with (
            mock.patch.object(dc, "get_converter", return_value=conv),
        ):
            summary = fill_pipeline.run_pipeline(
                str(self.xlsx),
                str(self.before_dir),
                str(self.out_dir),
                template=str(self.template),
            )
        # A1 成功写回；B1（skipped）不得产出，也不得抛异常中断整批
        self.assertTrue((self.out_dir / "A1.dxf").is_file())
        self.assertFalse((self.out_dir / "B1.dxf").is_file())
        self.assertEqual(summary["skipped"], ["B1"])
        self.assertEqual(summary["ok"], 1)


class FillPipelineMetaPriorityTest(unittest.TestCase):
    """模板库只存 meta JSON（原文件不存在）时，run_pipeline 直接读 meta 运行。

    回归：meta 优先分支不得再要求模板文件存在，也不得调用 template_to_dxf。
    """

    def setUp(self) -> None:
        import ezdxf
        import openpyxl

        from cadbatchassistant.core.common.template_meta import save_template_meta

        self.tmp = Path(tempfile.mkdtemp(prefix="fill_pipe_meta_"))
        self.before_dir = self.tmp / "input"
        self.out_dir = self.tmp / "output"
        self.before_dir.mkdir()
        self.out_dir.mkdir()
        # before DXF 含匹配文字 D-001
        doc_a = ezdxf.new("R2013")
        doc_a.modelspace().add_text(
            "D-001", dxfattribs={"insert": (10, 10, 0), "height": 3.0}
        )
        doc_a.saveas(self.before_dir / "A1.dxf")
        # 模板虚拟路径：文件不存在（模板库只存占位符 JSON）
        self.template = self.tmp / "templates" / "fill" / "tpl.dxf"
        self.template.parent.mkdir(parents=True)
        save_template_meta(
            self.template,
            {
                "placeholders": [
                    {
                        "text": "图号", "layer": "0",
                        "x": 10.0, "y": 20.0, "height": 3.0, "style": "",
                        "halign": 0, "valign": 0, "ref_text": "",
                        "entity_desc": {
                            "dxftype": "TEXT",
                            "attribs": {
                                "layer": "0",
                                "insert": (10.0, 20.0, 0.0),
                                "height": 3.0,
                                "style": "",
                                "halign": 0,
                                "valign": 0,
                            },
                            "layer_attribs": None,
                            "style_attribs": None,
                        },
                    }
                ]
            },
        )
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["图纸", "图号"])
        ws.append(["A1", "D-001"])
        self.xlsx = self.tmp / "data.xlsx"
        wb.save(self.xlsx)
        wb.close()

    def tearDown(self) -> None:
        import shutil

        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_run_with_meta_and_no_template_file(self) -> None:
        """meta 存在 + 模板文件不存在：正常运行，不调用 template_to_dxf。"""
        from unittest import mock

        from cadbatchassistant.core import dwg_converter as dc
        from cadbatchassistant.core.fill import fill_pipeline

        self.assertFalse(self.template.is_file())  # 模板库只有 meta JSON
        conv = mock.Mock()
        conv.resolve.return_value = ""
        conv.require_for_dwg.return_value = None
        conv.require_for_dwg.return_value = None
        with (
            mock.patch.object(dc, "get_converter", return_value=conv),
        ):
            summary = fill_pipeline.run_pipeline(
                str(self.xlsx),
                str(self.before_dir),
                str(self.out_dir),
                template=str(self.template),
            )
        self.assertTrue((self.out_dir / "A1.dxf").is_file())
        self.assertEqual(summary["ok"], 1)
        conv.template_to_dxf.assert_not_called()

    def test_run_meta_empty_placeholders_raises(self) -> None:
        """meta 的 placeholders 为空（手改 JSON）→ 报配置损坏，而非静默输出原图。"""
        from unittest import mock

        from cadbatchassistant.core import dwg_converter as dc
        from cadbatchassistant.core.common.template_meta import save_template_meta
        from cadbatchassistant.core.fill import fill_pipeline

        save_template_meta(self.template, {"placeholders": []})
        conv = mock.Mock()
        conv.resolve.return_value = ""
        conv.require_for_dwg.return_value = None
        conv.require_for_dwg.return_value = None
        with (
            mock.patch.object(dc, "get_converter", return_value=conv),
            self.assertRaises(ValueError) as ctx,
        ):
            fill_pipeline.run_pipeline(
                str(self.xlsx),
                str(self.before_dir),
                str(self.out_dir),
                template=str(self.template),
            )
        self.assertIn("占位配置损坏或为空", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
